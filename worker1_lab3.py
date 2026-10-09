import os
import random
import uuid
import zlib
from datetime import date

from openpyxl import load_workbook

from worker1 import (
    DAYS,
    INITIALS,
    PARAMETERS,
    PROJECTS,
    SERUMS,
    _self_check,
    build_workbook,
    log_created,
    results_dir,
)

# Hand-typed names. Worker 2 maps these back onto the Lab 1 lists.
SERUM_OUT = {
    "Pooled Normal Human Serum": "Pooled Normal Human Serun",
    "Individual Donor Serum": "individual donor serum",
    "Heat-Inactivated Human Serum": "heat inactivated human serum",
    "Fetal Bovine Serum": "fetal bovine serum",
}


def plate_seed(project, serum, parameter, initials, day):
    key = f"east|{project}|{serum}|{parameter}|{initials}|{day}"
    return zlib.adler32(key.encode()) & 0x7FFFFFFF


def _euro(value):
    return f"{float(value):.3f}".replace(".", ",")


def dirty(ws, _project, serum, parameter, initials):
    ws["A2"] = SERUM_OUT[serum]
    ws["A4"] = initials + " "
    ws["C1"] = "East Lab"
    ws["C2"] = "hand entry"
    today = date.today()
    ws["C3"] = f"{today.day} {today.strftime('%b %Y')}"
    ws["C4"] = "1:100"
    ws["A14"] = "edge wells look high, left the blanks"
    ws.cell(row=5, column=1, value="a")
    # A1:A2 and B1 typed with a decimal comma. A3 carries a unit. A4 is #N/A.
    for row, col in ((5, 2), (5, 3), (6, 2)):
        cell = ws.cell(row=row, column=col)
        cell.value = _euro(cell.value)
    suffix = "ng/mL" if parameter == "Concentration (ng/mL)" else "OD"
    cell = ws.cell(row=5, column=4)
    cell.value = f"{float(cell.value):.3f} {suffix}"
    ws.cell(row=5, column=5, value="#N/A")
    ws.cell(row=5, column=6, value=-0.015)
    ws.cell(row=5, column=7, value=-0.008)
    if parameter == "Percent Inhibition":
        ws.cell(row=5, column=8, value="45%")


def main():
    _self_check()
    project = random.choice(PROJECTS)
    serum = random.choice(SERUMS)
    parameter = random.choice(PARAMETERS)
    initials = random.choice(INITIALS)
    day = random.choice(DAYS)
    seed = plate_seed(project, serum, parameter, initials, day)
    low, high = 0.050, 2.500
    wb = build_workbook(
        project,
        serum,
        parameter,
        initials,
        seed=seed,
        low=low,
        high=high,
        day=day,
    )
    dirty(wb.active, project, serum, parameter, initials)
    today = date.today()
    suffix = uuid.uuid4().hex[:8]
    name = (
        f"final_plate {initials} {today.year}-{today.month:02d}-{today.day:02d}"
        f" {suffix}.xlsx"
    )
    path = os.path.join(results_dir("lab3"), name)
    wb.save(path)
    ws = load_workbook(path).active
    assert ws["C1"].value == "East Lab"
    assert ws["A2"].value == SERUM_OUT[serum]
    assert ws["A4"].value == initials + " "
    assert ws["C4"].value == "1:100"
    assert ws["A14"].value == "edge wells look high, left the blanks"
    assert ws.cell(row=5, column=1).value == "a"
    assert ws.cell(row=5, column=5).value == "#N/A"
    assert "," in str(ws.cell(row=5, column=2).value)
    log_created(path, project, serum, parameter, initials, day, seed, low, high)
    print(path)


if __name__ == "__main__":
    main()
