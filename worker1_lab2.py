import os
import random
import uuid
import zlib
from datetime import date

from openpyxl import load_workbook

from worker1 import (
    DAYS,
    INITIALS,
    PARAMETERS,
    PROJECTS,
    SERUMS,
    _self_check,
    build_workbook,
    log_created,
    results_dir,
)

# Plate-reader shorthand. Worker 2 maps these back onto the Lab 1 lists.
SERUM_OUT = {
    "Pooled Normal Human Serum": "Pooled NHS ",
    "Individual Donor Serum": "Donor serum ",
    "Heat-Inactivated Human Serum": "HI Human Serum ",
    "Fetal Bovine Serum": "FBS ",
}
PARAMETER_OUT = {
    "Absorbance (OD450)": "OD450",
    "Absorbance (OD405)": "OD405",
    "Absorbance (OD492)": "OD492",
    "Fluorescence (RFU)": "RFU",
    "Luminescence (RLU)": "RLU",
    "Concentration (ng/mL)": "ng/ml",
    "Percent Inhibition": "% inhibition",
}
RANGES = {
    "Absorbance (OD450)": (0.050, 2.500),
    "Absorbance (OD405)": (0.050, 2.500),
    "Absorbance (OD492)": (0.050, 2.500),
    "Fluorescence (RFU)": (100.0, 50_000.0),
    "Luminescence (RLU)": (1_000.0, 1_000_000.0),
    "Concentration (ng/mL)": (0.1, 500.0),
    "Percent Inhibition": (0.0, 100.0),
}


def plate_seed(project, serum, parameter, initials, day):
    key = f"west|{project}|{serum}|{parameter}|{initials}|{day}"
    return zlib.adler32(key.encode()) & 0x7FFFFFFF


def dirty(ws, project, serum, parameter, initials):
    ws["A1"] = project.replace("-", " ")
    ws["A2"] = SERUM_OUT[serum]
    ws["A3"] = PARAMETER_OUT[parameter]
    ws["A4"] = initials.lower()
    ws["C1"] = "West Lab"
    ws["C2"] = "BioTek Synergy H1"
    ws["C3"] = date.today().strftime("%m/%d/%Y")
    ws["C4"] = "endpoint"
    # Row 5 is plate row A. A1 saturated, A2 missed, A3:A5 exported as text.
    ws.cell(row=5, column=2, value="OVRFLW")
    # cell(value=None) means "leave it"; assign on the cell so the read is blank.
    ws.cell(row=5, column=3).value = None
    for col in (4, 5, 6):
        cell = ws.cell(row=5, column=col)
        cell.value = f"{float(cell.value):.3f}"


def main():
    _self_check()
    project = random.choice(PROJECTS)
    serum = random.choice(SERUMS)
    parameter = random.choice(PARAMETERS)
    initials = random.choice(INITIALS)
    day = random.choice(DAYS)
    seed = plate_seed(project, serum, parameter, initials, day)
    low, high = RANGES[parameter]
    wb = build_workbook(
        project,
        serum,
        parameter,
        initials,
        seed=seed,
        low=low,
        high=high,
        day=day,
    )
    dirty(wb.active, project, serum, parameter, initials)
    today = date.today()
    suffix = uuid.uuid4().hex[:8]
    name = (
        f"Synergy_{today.month:02d}{today.day:02d}{today.year}"
        f"_{initials.lower()}_{suffix}.xlsx"
    )
    path = os.path.join(results_dir("lab2"), name)
    wb.save(path)
    ws = load_workbook(path).active
    assert ws["C1"].value == "West Lab"
    assert ws["A1"].value == project.replace("-", " ")
    assert ws["A2"].value == SERUM_OUT[serum]
    assert ws["A3"].value == PARAMETER_OUT[parameter]
    assert ws.cell(row=5, column=2).value == "OVRFLW"
    assert ws.cell(row=5, column=3).value is None
    log_created(path, project, serum, parameter, initials, day, seed, low, high)
    print(path)


if __name__ == "__main__":
    main()
