import json
import os
import random
import subprocess
import sys
import tempfile
import uuid
from datetime import date, datetime

import openpyxl
from openpyxl import Workbook, load_workbook

_pool = random.Random(1)


def _unique(draw, n):
    seen = set()
    while len(seen) < n:
        seen.add(draw())
    return sorted(seen)


#A1
PROJECTS = _unique(lambda: f"Domino-{_pool.randint(0, 9_999_999):07d}", 7)
#A2
SERUMS = [
    "Pooled Normal Human Serum",
    "Individual Donor Serum",
    "Heat-Inactivated Human Serum",
    "Fetal Bovine Serum",
]
#A3
PARAMETERS = [
    "Absorbance (OD450)",
    "Absorbance (OD405)",
    "Absorbance (OD492)",
    "Fluorescence (RFU)",
    "Luminescence (RLU)",
    "Concentration (ng/mL)",
    "Percent Inhibition",
]
#A4
INITIALS = _unique(
    lambda: "".join(_pool.choice("ABCDEFGHIJKLMNOPQRSTUVWXYZ") for _ in range(2)),
    10,
)

assert len(PROJECTS) == 7
assert len(SERUMS) == 4
assert len(PARAMETERS) == 7
assert len(INITIALS) == 10

# Member 1 is A. A team of six would be A-F; this roster is 10, so A-J.
DAYS = ("day 1", "day 2", "day 3")


def analyst_letter(initials):
    return chr(ord("A") + INITIALS.index(initials))


#Standard 96 well plate
ROWS = "ABCDEFGH"
N_COLS = 12


def results_dir(subdir="results"):

    #Write under the dataset mount. Database comes later.

    dataset = "/mnt/data/Worker1Raw"
    if not os.path.isdir(dataset):
        raise FileNotFoundError(f"dataset directory missing: {dataset}")
    path = os.path.join(dataset, subdir)
    os.makedirs(path, exist_ok=True)
    return path


def code_version():
    try:
        out = subprocess.run(
            ["git", "rev-parse", "HEAD"],
            capture_output=True,
            text=True,
            timeout=2,
            cwd=os.path.dirname(os.path.abspath(__file__)),
        )
        if out.returncode == 0 and out.stdout.strip():
            return out.stdout.strip()
    except (OSError, subprocess.TimeoutExpired):
        pass
    return os.environ.get("FLYTE_INTERNAL_TASK_VERSION") or "unknown"


def environment():
    py = f"{sys.version_info.major}.{sys.version_info.minor}.{sys.version_info.micro}"
    hardware = os.environ.get("DOMINO_HARDWARE_TIER_ID") or "local"
    return f"python {py}; openpyxl {openpyxl.__version__}; hardware {hardware}"


def run_record(path, project, serum, parameter, initials, day, seed, low, high, when=None):
    return {
        "when": when or datetime.now().isoformat(timespec="seconds"),
        "what": os.path.basename(path),
        "where": path,
        "lab": os.path.basename(os.path.dirname(path)),
        "input_id": f"{project}|{serum}|{parameter}|{initials}|{day}",
        "code_version": code_version(),
        "parameters": {"seed": seed, "low": low, "high": high},
        "environment": environment(),
        "run_id": os.environ.get("DOMINO_RUN_ID") or "local",
    }


def _append(path, line):
    os.makedirs(os.path.dirname(path) or ".", exist_ok=True)
    with open(path, "a", encoding="utf-8") as handle:
        handle.write(line)
        handle.flush()
        os.fsync(handle.fileno())


def _log_run_id():
    # ponytail: one file per run. A shared worker1.txt is replaced by whichever job finishes last.
    return os.environ.get("DOMINO_RUN_ID") or f"local-{uuid.uuid4().hex[:8]}"


def log_destinations(plate_path):
    # Each job has its own /mnt/artifacts. The dataset file is the one all three jobs share.
    plate_dir = os.path.dirname(plate_path)
    dataset = os.path.dirname(plate_dir)
    dests = [
        os.path.join(dataset, "worker1", f"{_log_run_id()}.txt"),
        os.path.join(plate_dir, "worker1.txt"),
    ]
    root = os.environ.get("DOMINO_ARTIFACTS_DIR") or "/mnt/artifacts"
    dests.append(os.path.join(root, "worker1.txt"))
    return dests


def log_created(path, project, serum, parameter, initials, day, seed, low, high):
    record = run_record(path, project, serum, parameter, initials, day, seed, low, high)
    line = json.dumps(record, separators=(",", ":")) + "\n"
    for dest in log_destinations(path):
        _append(dest, line)
    print(line, end="", flush=True)


def fake_readings(seed=4, low=0.050, high=2.500):

    #N_Cols AND ROWS
    #Are standard and defined up top
    #Uniform Range Choosen because
    #Usual readable window for an ELISA plate reader at 450 nm
    rng = random.Random(seed)
    return [
        [round(rng.uniform(low, high), 3) for _ in range(N_COLS)]
        for _ in ROWS
    ]


def build_workbook(project, serum, parameter, initials, seed=4, low=0.050, high=2.500, day="day 1"):
    wb = Workbook()
    ws = wb.active
    ws.title = "Plate"
    ws["A1"] = project
    ws["A2"] = serum
    ws["A3"] = parameter
    ws["A4"] = initials
    ws["B1"] = day
    ws["B2"] = analyst_letter(initials)
    ws["C1"] = "North Lab"
    ws["C3"] = date.today().strftime("%d%b%Y")



    for col in range(1, N_COLS + 1):
        ws.cell(row=5, column=col + 1, value=col)


    readings = fake_readings(seed, low, high)

    
    for r, letter in enumerate(ROWS):
        ws.cell(row=5 + r, column=1, value=letter)
        for c, value in enumerate(readings[r]):
            ws.cell(row=5 + r, column=2 + c, value=value)
    return wb


def assert_layout(path, project, serum, parameter, initials, day="day 1"):
    ws = load_workbook(path).active
    assert ws["A1"].value == project
    assert ws["A2"].value == serum
    assert ws["A3"].value == parameter
    assert ws["A4"].value == initials
    assert ws["B1"].value == day
    assert ws["B2"].value == analyst_letter(initials)
    assert ws["C1"].value == "North Lab"
    assert ws["C3"].value == date.today().strftime("%d%b%Y")
    cells = [
        cell.value
        for row in ws.iter_rows(min_row=5, max_row=12, min_col=2, max_col=13)
        for cell in row
    ]
    assert len(cells) == 96
    assert all(isinstance(value, (int, float)) for value in cells)


def _self_check():
    record = run_record(
        "/mnt/data/Worker1Raw/lab1/assay.xlsx",
        "Domino-0000001",
        "Pooled Normal Human Serum",
        "Absorbance (OD450)",
        "AB",
        "day 1",
        4,
        0.050,
        2.500,
        when="2026-10-09T10:19:00",
    )
    for key in (
        "when",
        "what",
        "where",
        "input_id",
        "code_version",
        "parameters",
        "environment",
        "run_id",
    ):
        assert record.get(key) not in (None, "")
    assert record["input_id"] == (
        "Domino-0000001|Pooled Normal Human Serum|Absorbance (OD450)|AB|day 1"
    )
    assert record["parameters"] == {"seed": 4, "low": 0.050, "high": 2.500}
    assert record["what"] == "assay.xlsx"
    assert record["where"] == "/mnt/data/Worker1Raw/lab1/assay.xlsx"
    root = tempfile.mkdtemp()
    plate = os.path.join(root, "lab2", "assay.xlsx")
    saved_artifacts = os.environ.get("DOMINO_ARTIFACTS_DIR")
    saved_run = os.environ.get("DOMINO_RUN_ID")
    os.environ["DOMINO_ARTIFACTS_DIR"] = os.path.join(root, "missing-artifacts")
    os.environ["DOMINO_RUN_ID"] = "run-check"
    try:
        dests = log_destinations(plate)
    finally:
        if saved_artifacts is None:
            os.environ.pop("DOMINO_ARTIFACTS_DIR", None)
        else:
            os.environ["DOMINO_ARTIFACTS_DIR"] = saved_artifacts
        if saved_run is None:
            os.environ.pop("DOMINO_RUN_ID", None)
        else:
            os.environ["DOMINO_RUN_ID"] = saved_run
    assert dests == [
        os.path.join(root, "worker1", "run-check.txt"),
        os.path.join(root, "lab2", "worker1.txt"),
        os.path.join(root, "missing-artifacts", "worker1.txt"),
    ]
    line = json.dumps(record, separators=(",", ":")) + "\n"
    for dest in dests:
        _append(dest, line)
        assert open(dest, encoding="utf-8").read().endswith(line)


def main():
    _self_check()
    project = random.choice(PROJECTS)
    serum = random.choice(SERUMS)
    parameter = random.choice(PARAMETERS)
    initials = random.choice(INITIALS)
    day = random.choice(DAYS)
    seed, low, high = 4, 0.050, 2.500
    today = date.today()
    suffix = uuid.uuid4().hex[:8]
    name = (
        f"assay_plate_{today.year}-{today.month:02d}-{today.day:02d}"
        f"_{initials}_{suffix}.xlsx"
    )
    path = os.path.join(results_dir("lab1"), name)
    build_workbook(
        project, serum, parameter, initials, seed=seed, low=low, high=high, day=day
    ).save(path)

    #Checkworkbook
    assert_layout(path, project, serum, parameter, initials, day)
    log_created(path, project, serum, parameter, initials, day, seed, low, high)
    print(path)


if __name__ == "__main__":
    main()
