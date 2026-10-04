import os
import random
from datetime import date

from openpyxl import Workbook, load_workbook


#A1
PROJECT_NAME = "Domino-1234567"
#A2
SERUM_TYPE = "Pooled Normal Human Serum"
#A3
PARAMETER = "Absorbance (OD450)"
#A4
Initals = "GM"

#Standard 96 well plate
ROWS = "ABCDEFGH"
N_COLS = 12


def results_dir():

    #Write under the dataset mount. Database comes later.

    dataset = "/domino/datasets/local/Domino_Takehome"
    if not os.path.isdir(dataset):
        raise FileNotFoundError(f"dataset directory missing: {dataset}")
    path = os.path.join(dataset, "results")
    os.makedirs(path, exist_ok=True)
    return path


def fake_readings():

    #N_Cols AND ROWS 
    #Are standard and defined up top


    random.seed(4)

    #Uniform Range Choosen because
    #Usual readable window for an ELISA plate reader at 450 nm
    return [
        [round(random.uniform(0.050, 2.500), 3) for _ in range(N_COLS)]
        for _ in ROWS
    ]


def build_workbook():
    wb = Workbook()
    ws = wb.active
    ws.title = "Plate"
    ws["A1"] = PROJECT_NAME
    ws["A2"] = SERUM_TYPE
    ws["A3"] = PARAMETER
    ws["A4"] = Initals



    for col in range(1, N_COLS + 1):
        ws.cell(row=5, column=col + 1, value=col)


    readings = fake_readings()

    
    for r, letter in enumerate(ROWS):
        ws.cell(row=5 + r, column=1, value=letter)
        for c, value in enumerate(readings[r]):
            ws.cell(row=5 + r, column=2 + c, value=value)
    return wb


def assert_layout(path):
    ws = load_workbook(path).active
    assert ws["A1"].value == PROJECT_NAME
    assert ws["A2"].value == SERUM_TYPE
    assert ws["A3"].value == PARAMETER
    assert ws["A4"].value == Initals
    cells = [
        cell.value
        for row in ws.iter_rows(min_row=5, max_row=12, min_col=2, max_col=13)
        for cell in row
    ]
    assert len(cells) == 96
    assert all(isinstance(value, (int, float)) for value in cells)


def main():
    today = date.today()
    name = f"assay_plate_{today.year}-{today.month:02d}-{today.day:02d}_{Initals}.xlsx"
    path = os.path.join(results_dir(), name)
    build_workbook().save(path)

    #Checkworkbook
    assert_layout(path)
    print(path)


if __name__ == "__main__":
    main()
