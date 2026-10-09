import os
import re
import sqlite3
from datetime import date, datetime

from openpyxl import load_workbook

import db

from worker1 import (
    INITIALS,
    N_COLS,
    PARAMETERS,
    PROJECTS,
    ROWS,
    SERUMS,
    analyst_letter,
    build_workbook,
)

# ponytail: a comma is always a decimal mark; "1.234,5" is rejected.
OVERFLOW = {"OVRFLW", "OVERFLOW", "****"}
MISSING = {"#N/A", "#DIV/0!", "ND", "N.A.", "N.A", "NA"}
SERUM_ALIASES = {
    "pooled normal human serum": "Pooled Normal Human Serum",
    "pooled nhs": "Pooled Normal Human Serum",
    "pooled normal human serun": "Pooled Normal Human Serum",
    "individual donor serum": "Individual Donor Serum",
    "donor serum": "Individual Donor Serum",
    "heat-inactivated human serum": "Heat-Inactivated Human Serum",
    "heat inactivated human serum": "Heat-Inactivated Human Serum",
    "hi human serum": "Heat-Inactivated Human Serum",
    "fetal bovine serum": "Fetal Bovine Serum",
    "fbs": "Fetal Bovine Serum",
}
PARAMETER_ALIASES = {
    "absorbance (od450)": "Absorbance (OD450)",
    "od450": "Absorbance (OD450)",
    "od 450": "Absorbance (OD450)",
    "absorbance (od405)": "Absorbance (OD405)",
    "od405": "Absorbance (OD405)",
    "od 405": "Absorbance (OD405)",
    "absorbance (od492)": "Absorbance (OD492)",
    "od492": "Absorbance (OD492)",
    "od 492": "Absorbance (OD492)",
    "fluorescence (rfu)": "Fluorescence (RFU)",
    "rfu": "Fluorescence (RFU)",
    "luminescence (rlu)": "Luminescence (RLU)",
    "rlu": "Luminescence (RLU)",
    "concentration (ng/ml)": "Concentration (ng/mL)",
    "ng/ml": "Concentration (ng/mL)",
    "percent inhibition": "Percent Inhibition",
    "% inhibition": "Percent Inhibition",
    "%inhib": "Percent Inhibition",
}
LABS = {
    "lab1": "Lab1",
    "north lab": "Lab1",
    "west lab": "Lab2",
    "lab2": "Lab2",
    "east lab": "Lab3",
    "lab3": "Lab3",
}

RAW = "/mnt/data/Worker1Raw"
artifacts_logs = "/mnt/artifacts/worker2.txt"


def _token(value):
    return "" if value is None else " ".join(str(value).strip().lower().split())


def canon_project(value):
    text = "" if value is None else str(value).strip()
    match = re.fullmatch(r"(?i)domino[-\s]?(\d{7})", text)
    if not match:
        return text
    return f"Domino-{int(match.group(1)):07d}"


def canon_name(value, aliases, allowed):
    text = "" if value is None else str(value).strip()
    if text in allowed:
        return text
    return aliases.get(_token(value), text)


def canon_lab(value):
    if value is None or str(value).strip() == "":
        return ""
    return LABS.get(_token(value), str(value).strip())


def _stamp(value):
    return value.strftime("%d%b%Y")


def canon_day(value):
    text = "" if value is None else " ".join(str(value).strip().lower().split())
    match = re.fullmatch(r"day (\d+)", text)
    if not match:
        return ""
    number = int(match.group(1))
    if number < 1:
        return ""
    return f"day {number}"


def canon_date(value):
    text = "" if value is None else str(value).strip()
    if not text:
        return ""
    for fmt in ("%m/%d/%Y", "%d %b %Y", "%Y-%m-%d", "%d%b%Y"):
        try:
            return _stamp(datetime.strptime(text, fmt).date())
        except ValueError:
            continue
    return ""


def date_from_name(filename):
    text = os.path.basename(filename or "")
    match = re.search(r"(20\d{2})-(\d{2})-(\d{2})", text)
    if match:
        try:
            return _stamp(date(int(match.group(1)), int(match.group(2)), int(match.group(3))))
        except ValueError:
            return ""
    match = re.search(r"(?<!\d)(\d{2})(\d{2})(20\d{2})(?!\d)", text)
    if not match:
        return ""
    try:
        return _stamp(date(int(match.group(3)), int(match.group(1)), int(match.group(2))))
    except ValueError:
        return ""


def parse_well(value, well):
    if isinstance(value, bool):
        raise ValueError(f"skipped: well {well} is not a number")
    if isinstance(value, (int, float)):
        return float(value), ""
    if value is None:
        return None, "blank"
    if not isinstance(value, str):
        raise ValueError(f"skipped: well {well} is not a number")
    text = value.strip()
    token = text.upper()
    if text == "" or token == "BLANK":
        return None, "blank"
    if token in OVERFLOW:
        return None, "overflow"
    if token in MISSING:
        return None, "na"
    if text.endswith("%"):
        text = text[:-1].strip()
    lower = text.lower()
    for unit in ("ng/ml", "rfu", "rlu", "od"):
        if lower.endswith(unit):
            text = text[: -len(unit)].strip()
            break
    if "," in text and "." in text:
        raise ValueError(f"skipped: well {well} is not a number")
    try:
        return float(text.replace(",", ".")), ""
    except ValueError:
        raise ValueError(f"skipped: well {well} is not a number") from None


def plate_rows(ws, filename=""):
    project = canon_project(ws["A1"].value)
    serum = canon_name(ws["A2"].value, SERUM_ALIASES, SERUMS)
    parameter = canon_name(ws["A3"].value, PARAMETER_ALIASES, PARAMETERS)
    initials = "" if ws["A4"].value is None else str(ws["A4"].value).strip().upper()
    lab = canon_lab(ws["C1"].value)
    instrument = "" if ws["C2"].value is None else str(ws["C2"].value).strip()
    run_date = canon_date(ws["C3"].value) or date_from_name(filename)
    fields = (
        ("A1", "project", PROJECTS, project),
        ("A2", "serum", SERUMS, serum),
        ("A3", "parameter", PARAMETERS, parameter),
        ("A4", "initials", INITIALS, initials),
    )
    for cell, label, allowed, value in fields:
        if value not in allowed:
            raise ValueError(f'skipped: {cell} {label} not allowed: "{value}"')
    raw_day = ws["B1"].value
    day = canon_day(raw_day)
    if not day:
        raise ValueError(f'skipped: B1 day not allowed: "{raw_day}"')
    analyst = "" if ws["B2"].value is None else str(ws["B2"].value).strip().upper()
    if analyst != analyst_letter(initials):
        raise ValueError(f'skipped: B2 analyst not allowed: "{analyst}"')
    for label, value in (("lab", lab), ("instrument", instrument), ("run_date", run_date)):
        if not value:
            raise ValueError(f"skipped: plate missing {label}")
    rows = []
    flagged = 0
    for r, letter in enumerate(ROWS):
        got = ws.cell(row=5 + r, column=1).value
        got = "" if got is None else str(got).strip().upper()
        if got != letter:
            raise ValueError(f"skipped: row {letter} label is {got!r}")
        for c in range(1, N_COLS + 1):
            well = f"{letter}{c}"
            number, flag = parse_well(ws.cell(row=5 + r, column=c + 1).value, well)
            if flag:
                flagged += 1
            rows.append(
                [
                    project,
                    serum,
                    parameter,
                    initials,
                    well,
                    "" if number is None else f"{number:.3f}",
                    lab,
                    instrument,
                    run_date,
                    flag,
                    day,
                    analyst,
                ]
            )
    return (project, serum, parameter, initials, lab, day, run_date), rows, flagged


def xlsx_files(root):
    found = []
    for dirpath, _, names in os.walk(root):
        for name in names:
            if name.lower().endswith(".xlsx"):
                found.append(os.path.join(dirpath, name))
    return sorted(found)


def log_line(filename, quartet, message):
    stamp = datetime.now().isoformat(timespec="seconds")
    plate = ", ".join(map(str, quartet)) if quartet else ""
    os.makedirs(os.path.dirname(artifacts_logs) or ".", exist_ok=True)
    with open(artifacts_logs, "a", encoding="utf-8") as handle:
        handle.write(f"{stamp} {filename}: {plate} | {message}\n")


def log_skip(filename, quartet, reason):
    if not reason.startswith("skipped:"):
        reason = f"skipped: {reason}"
    log_line(filename, quartet, reason)
    return reason


def _self_check():
    global artifacts_logs
    from worker1_lab2 import dirty as dirty_west
    from worker1_lab3 import dirty as dirty_east

    today = date.today().strftime("%d%b%Y")
    ws = build_workbook(PROJECTS[0], SERUMS[0], PARAMETERS[0], INITIALS[0]).active
    ws["C2"] = "SpectraMax iD3"
    quartet, rows, flagged = plate_rows(ws, f"assay_plate_{date.today().isoformat()}_AB_abcd.xlsx")
    assert flagged == 0 and len(rows) == 96 and quartet[4:] == ("Lab1", "day 1", today)
    assert rows[0][7:] == ["SpectraMax iD3", today, "", "day 1", "A"]
    missing = build_workbook(PROJECTS[0], SERUMS[0], PARAMETERS[0], INITIALS[0]).active
    try:
        plate_rows(missing, "nodate.xlsx")
        raise AssertionError("missing instrument accepted")
    except ValueError as exc:
        assert str(exc) == "skipped: plate missing instrument"
    missing["C2"] = "SpectraMax iD3"
    missing["C3"] = None
    try:
        plate_rows(missing, "nodate.xlsx")
        raise AssertionError("missing run_date accepted")
    except ValueError as exc:
        assert str(exc) == "skipped: plate missing run_date"
    missing["C3"] = today
    missing["C1"] = None
    try:
        plate_rows(missing, "nodate.xlsx")
        raise AssertionError("missing lab accepted")
    except ValueError as exc:
        assert str(exc) == "skipped: plate missing lab"
    assert analyst_letter(INITIALS[0]) == "A"
    assert analyst_letter(INITIALS[1]) == "B"
    ws["C3"] = "12Dec2023"
    _, stamped, _ = plate_rows(ws)
    assert stamped[0][8] == "12Dec2023"
    ws["C3"] = None
    _, named, _ = plate_rows(ws, "assay_plate_2023-12-12_AB.xlsx")
    assert named[0][8] == "12Dec2023"

    west = build_workbook(PROJECTS[0], SERUMS[3], PARAMETERS[0], INITIALS[0]).active
    original_a3 = west.cell(row=5, column=4).value
    dirty_west(west, PROJECTS[0], SERUMS[3], PARAMETERS[0], INITIALS[0])
    string_well = west.cell(row=5, column=4).value
    assert isinstance(string_well, str)
    quartet, rows, flagged = plate_rows(west)
    by_well = {row[4]: (row[5], row[9]) for row in rows}
    assert quartet == (PROJECTS[0], SERUMS[3], PARAMETERS[0], INITIALS[0], "Lab2", "day 1", today)
    assert flagged == 2 and len(rows) == 96
    assert by_well["A1"] == ("", "overflow") and by_well["A2"] == ("", "blank")
    assert by_well["A3"] == (f"{float(string_well):.3f}", "")
    assert by_well["A3"][0] == f"{original_a3:.3f}"
    assert rows[0][7] == "BioTek Synergy H1" and rows[0][8] == today

    percent = PARAMETERS[-1]
    east = build_workbook(PROJECTS[1], SERUMS[2], percent, INITIALS[1]).active
    original_a1 = east.cell(row=5, column=2).value
    original_a3 = east.cell(row=5, column=4).value
    dirty_east(east, PROJECTS[1], SERUMS[2], percent, INITIALS[1])
    assert east.cell(row=5, column=1).value == "a"
    assert east["A14"].value
    quartet, rows, flagged = plate_rows(east)
    by_well = {row[4]: (row[5], row[9]) for row in rows}
    assert quartet == (PROJECTS[1], SERUMS[2], percent, INITIALS[1], "Lab3", "day 1", today)
    assert flagged == 1 and len(rows) == 96
    assert by_well["A4"] == ("", "na")
    assert by_well["A1"][0] == f"{original_a1:.3f}" and by_well["A1"][1] == ""
    assert by_well["A3"][0] == f"{original_a3:.3f}"
    assert by_well["A5"] == ("-0.015", "") and by_well["A6"] == ("-0.008", "")
    assert by_well["A7"] == ("45.000", "")
    assert rows[0][7] == "hand entry" and rows[0][8] == today
    assert rows[0][10:] == ["day 1", "B"]

    saved = artifacts_logs
    folder = os.path.join(os.path.dirname(os.path.abspath(__file__)), "_w2_log_check")
    artifacts_logs = os.path.join(folder, "worker2.txt")
    os.makedirs(folder, exist_ok=True)
    conn = None
    try:
        log_skip("bad.xlsx", quartet, "skipped: A2 serum not allowed")
        log_line("good.xlsx", quartet, "accepted")
        logged = open(artifacts_logs, encoding="utf-8").read()
        assert f"{', '.join(quartet)} | skipped: A2 serum not allowed" in logged
        assert "good.xlsx:" in logged and f"{', '.join(quartet)} | accepted" in logged
        db_path = os.path.join(folder, "plates.sqlite")
        conn = db.ensure(db_path)
        empty, columns = db.load_rows(db_path)
        assert empty == [] and columns == list(db.VIEW_COLUMNS)
        db.insert_plate(conn, quartet, rows, "good.xlsx")
        db.log_outcome(conn, "good.xlsx", "accepted", "accepted", quartet)
        conn.commit()
        loaded, columns = db.load_rows(db_path)
        assert columns == list(db.VIEW_COLUMNS)
        by_cell = {row["cell"]: row for row in loaded}
        assert len(by_cell) == 96
        assert by_cell["A4"]["value"] == "" and by_cell["A4"]["flag"] == "na"
        assert by_cell["A5"]["value"] == "-0.015" and by_cell["A6"]["value"] == "-0.008"
        assert by_cell["A7"]["value"] == "45.000"
        head = by_cell["A1"]
        assert head["project"] == quartet[0]
        assert head["serums"] == quartet[1]
        assert head["PARAMETERS"] == quartet[2]
        assert head["initals"] == quartet[3]
        assert head["lab"] == "Lab3" and head["day"] == "day 1"
        assert head["run_date"] == quartet[6]
        assert head["analyst"] == "B" and head["instrument"] == "hand entry"
        try:
            db.insert_plate(conn, quartet, rows, "again.xlsx")
            raise AssertionError("duplicate plate accepted")
        except sqlite3.IntegrityError:
            conn.rollback()
        other = quartet[:5] + ("day 2", quartet[6])
        db.insert_plate(conn, other, rows, "day2.xlsx")
        conn.rollback()
        assert conn.execute("SELECT COUNT(*) FROM plates").fetchone()[0] == 1
        assert conn.execute("SELECT COUNT(*) FROM wells").fetchone()[0] == 96
        logged_row = conn.execute(
            "SELECT outcome, source_file, day, run_date FROM ingest_log"
        ).fetchone()
        assert logged_row["outcome"] == "accepted" and logged_row["source_file"] == "good.xlsx"
        assert logged_row["day"] == 1
        assert logged_row["run_date"] == datetime.strptime(quartet[6], "%d%b%Y").date().isoformat()
    finally:
        artifacts_logs = saved
        if conn is not None:
            conn.close()
        for name in os.listdir(folder):
            os.remove(os.path.join(folder, name))
        os.rmdir(folder)


def main():


    #Confirm a fake build of workbook.
    _self_check()

    if not os.path.isdir(RAW):
        raise FileNotFoundError(f"dataset directory missing: {RAW}")

    conn = db.ensure()
    accepted = skipped = 0

    #Cycle Through All Workbooks in the directory.
    try:
        for path in xlsx_files(RAW):
            name = os.path.basename(path)
            quartet = None
            try:
                ws = load_workbook(path).active
                quartet = tuple(ws[cell].value for cell in ("A1", "A2", "A3", "A4"))
                quartet, rows, flagged = plate_rows(ws, name)
            except Exception as exc:
                reason = log_skip(name, quartet, str(exc))
                db.log_outcome(conn, name, "skipped", reason, quartet)
                conn.commit()
                skipped += 1
                continue
            try:
                db.insert_plate(conn, quartet, rows, name)
            except sqlite3.IntegrityError as exc:
                conn.rollback()
                if "UNIQUE" not in str(exc):
                    raise
                reason = log_skip(
                    name,
                    quartet,
                    "skipped: plate already in mock_db (" + ", ".join(quartet) + ")",
                )
                db.log_outcome(conn, name, "skipped", reason, quartet)
                conn.commit()
                skipped += 1
                continue
            detail = "accepted" if not flagged else f"accepted, flagged {flagged} wells"
            db.log_outcome(conn, name, "accepted", detail, quartet)
            conn.commit()
            log_line(name, quartet, detail)
            accepted += 1
    finally:
        conn.close()
    print(f"accepted {accepted}, skipped {skipped}")


if __name__ == "__main__":
    main()
